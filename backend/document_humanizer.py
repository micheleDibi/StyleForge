"""
Umanizzazione di testi/documenti lunghi.

Due capacità:
  1. humanize_long_text(): umanizza un testo arbitrariamente lungo dividendolo in
     CHUNK (mai a metà frase), applicando per ogni chunk la riscrittura nello stile
     appreso + il pipeline anti-AI completo. Risolve il crash su testi lunghi.
  2. humanize_docx_inplace(): round-trip di un .docx mantenendo il TEMPLATE originale
     (frontespizio, indice, titoli, font, margini, numerazione). Sostituisce SOLO il
     testo dei paragrafi del corpo, in batch che PRESERVANO la corrispondenza 1:1 coi
     paragrafi (così i titoli non si spostano). Niente controlled_paraphrase qui:
     riordina/fonde le frasi e romperebbe la mappatura dei paragrafi; l'anti-AI è dato
     dalla riscrittura nello stile + dal pass algoritmico per-paragrafo (gratis).
"""

import re
import logging

from ai_exceptions import InsufficientCreditsError

logger = logging.getLogger(__name__)


# ═══════════════════════════════════════════════════════════════════════════
# CHUNKING (testo lungo)
# ═══════════════════════════════════════════════════════════════════════════
def split_into_chunks(text: str, target_words: int = 3000, max_words: int = 3500):
    """Divide il testo in chunk ~target_words, su confini di paragrafo; mai a metà
    frase. Un paragrafo più lungo di max_words viene spezzato su confini di frase."""
    paragraphs = re.split(r'\n\s*\n', text or '')
    chunks, buf, buf_words = [], [], 0
    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        pw = len(para.split())
        if pw > max_words:
            if buf:
                chunks.append("\n\n".join(buf)); buf, buf_words = [], 0
            chunks.extend(_split_paragraph_on_sentences(para, target_words))
            continue
        if buf and buf_words + pw > target_words:
            chunks.append("\n\n".join(buf)); buf, buf_words = [], 0
        buf.append(para); buf_words += pw
    if buf:
        chunks.append("\n\n".join(buf))
    return [c for c in chunks if c.strip()]


def _split_paragraph_on_sentences(para: str, target_words: int):
    sentences = re.split(r'(?<=[.!?])\s+', para)
    chunks, buf, buf_words = [], [], 0
    for s in sentences:
        sw = len(s.split())
        if buf and buf_words + sw > target_words:
            chunks.append(" ".join(buf)); buf, buf_words = [], 0
        buf.append(s); buf_words += sw
    if buf:
        chunks.append(" ".join(buf))
    return chunks


def humanize_long_text(text: str, session_client, profile: str = 'informal', progress_cb=None) -> str:
    """Umanizza un testo lungo a chunk, col pipeline anti-AI completo per chunk."""
    import config
    from anti_ai_pipeline import apply_anti_ai_pipeline

    chunks = split_into_chunks(text)
    if not chunks:
        return text
    out = []
    n = len(chunks)
    for i, chunk in enumerate(chunks):
        rewritten = session_client.umanizzazione_chunk(chunk, profile=profile)
        rewritten = apply_anti_ai_pipeline(
            rewritten,
            profile=profile,
            target_words=len(chunk.split()),
            paraphrase_enabled=getattr(config, 'ANTI_AI_PARAPHRASE_ENABLED', True),
            paraphrase_rounds=getattr(config, 'ANTI_AI_PARAPHRASE_ROUNDS', 2),
            paraphrase_model=getattr(config, 'ANTI_AI_PARAPHRASE_MODEL', None),
            rewrite_enabled=getattr(config, 'ANTI_AI_REWRITE_ENABLED', False),
            rewrite_model=getattr(config, 'ANTI_AI_REWRITE_MODEL', None),
            algo_enabled=getattr(config, 'ANTI_AI_ALGO_ENABLED', True),
            label=f'chunk{i + 1}',
        )
        out.append(rewritten)
        if progress_cb:
            try:
                progress_cb(int((i + 1) / n * 100))
            except Exception:  # noqa: BLE001
                pass
    return "\n\n".join(out)


# ═══════════════════════════════════════════════════════════════════════════
# DOCX ROUND-TRIP (mantiene il template originale)
# ═══════════════════════════════════════════════════════════════════════════
# Stili da NON umanizzare (titoli, indice, didascalie, bibliografia, citazioni).
_SKIP_STYLE_PREFIXES = (
    'Heading', 'Title', 'TOC', 'Caption', 'Subtitle',
    'Titolo', 'Sottotitolo', 'Didascalia', 'Indice', 'Sommario',
)
_SKIP_STYLE_NAMES = {
    'Bibliography', 'Bibliografia', 'Quote', 'Intense Quote', 'Citazione',
}


def is_body_paragraph(p) -> bool:
    """True se il paragrafo è testo del corpo da umanizzare (non titolo/indice/
    bibliografia/didascalia, non vuoto, almeno 4 parole)."""
    txt = (p.text or '').strip()
    if not txt:
        return False
    try:
        name = (p.style.name or '') if p.style else ''
    except Exception:  # noqa: BLE001
        name = ''
    if any(name.startswith(pre) for pre in _SKIP_STYLE_PREFIXES):
        return False
    if name in _SKIP_STYLE_NAMES:
        return False
    if len(txt.split()) < 4:
        return False
    return True


def set_paragraph_text(p, new_text: str) -> None:
    """Sostituisce il testo del paragrafo mantenendone lo STILE (riusa il primo run
    per la formattazione di carattere, elimina gli altri). La formattazione inline
    (es. una parola in grassetto a metà) va persa: accettabile, le parole cambiano."""
    runs = p.runs
    if runs:
        runs[0].text = new_text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(new_text)


_MARKER_RE = re.compile(r'<<<\s*S\s*\d+\s*>>>')


def _build_segments_payload(texts) -> str:
    return "\n".join(f"<<<S{i}>>>\n{t}" for i, t in enumerate(texts))


def _parse_segments(rewritten: str, expected: int):
    """Ridivide la risposta sui marcatori <<<Sk>>>. Ritorna la lista dei segmenti
    se il conteggio combacia e nessuno è vuoto, altrimenti None."""
    parts = _MARKER_RE.split(rewritten or '')
    segs = [p.strip() for p in parts[1:]]  # scarta il preambolo prima del 1° marcatore
    if len(segs) != expected:
        return None
    if any(not s for s in segs):
        return None
    return segs


def _clean_segment(s: str) -> str:
    s = _MARKER_RE.sub('', s or '')
    s = s.strip()
    s = re.sub(r'^\s*-{3,}\s*', '', s)
    s = re.sub(r'\s*-{3,}\s*$', '', s)
    return s.strip()


def _per_paragraph_fallback(texts, session_client, profile):
    """Fallback robusto: una riscrittura per paragrafo → corrispondenza 1:1 garantita."""
    out = []
    for t in texts:
        try:
            out.append(session_client.umanizzazione_chunk(t, profile=profile))
        except InsufficientCreditsError:
            raise
        except Exception as e:  # noqa: BLE001
            logger.warning(f"Fallback per-paragrafo fallito: {e}; tengo l'originale")
            out.append(t)
    return out


def humanize_docx_inplace(src_path, dst_path, session_client, profile: str = 'academic',
                          progress_cb=None, batch_words: int = 2500) -> str:
    """
    Umanizza un .docx mantenendo il template originale: sostituisce SOLO il testo dei
    paragrafi del corpo, in batch che preservano la corrispondenza 1:1. Ritorna dst_path.
    """
    from docx import Document as DocxDocument
    from anti_ai_processor import humanize_text_post_processing

    doc = DocxDocument(str(src_path))
    body = [p for p in doc.paragraphs if is_body_paragraph(p)]
    total = len(body)
    if total == 0:
        doc.save(str(dst_path))
        return str(dst_path)

    # Raggruppa paragrafi-corpo consecutivi fino a ~batch_words parole.
    batches, cur, cur_words = [], [], 0
    for p in body:
        w = len(p.text.split())
        if cur and cur_words + w > batch_words:
            batches.append(cur); cur, cur_words = [], 0
        cur.append(p); cur_words += w
    if cur:
        batches.append(cur)

    done = 0
    for batch in batches:
        texts = [p.text for p in batch]
        if len(texts) == 1:
            segs = _per_paragraph_fallback(texts, session_client, profile)
        else:
            joined = _build_segments_payload(texts)
            try:
                rewritten = session_client.umanizzazione_segments(joined, len(texts), profile)
                segs = _parse_segments(rewritten, len(texts))
            except InsufficientCreditsError:
                raise
            except Exception as e:  # noqa: BLE001
                logger.warning(f"Riscrittura a segmenti fallita: {e}; uso il fallback")
                segs = None
            if segs is None:
                segs = _per_paragraph_fallback(texts, session_client, profile)

        for p, seg in zip(batch, segs):
            seg = _clean_segment(seg)
            try:
                seg = humanize_text_post_processing(seg, profile='academic')
            except Exception as e:  # noqa: BLE001
                logger.warning(f"Pass algoritmico per-paragrafo fallito: {e}")
            if seg and seg.strip():
                set_paragraph_text(p, seg.strip())

        done += len(batch)
        if progress_cb:
            try:
                progress_cb(int(done / total * 100))
            except Exception:  # noqa: BLE001
                pass

    doc.save(str(dst_path))
    return str(dst_path)
